"""Chuyển file Excel cũ "Tổng hợp thông tin công việc Fairplay Sports.xlsx" thành CSV để nhập vào CRM.

Chạy:  python3 tools/convert_excel.py "<đường dẫn file xlsx>" import/leads.csv
Lấy 2 sheet: "DS KH ads FB" (lead cá nhân từ quảng cáo) và "DS Giải đấu" (deal doanh nghiệp).
File CSV chứa dữ liệu khách hàng thật — KHÔNG commit lên GitHub (thư mục import/ đã được .gitignore).
"""
import csv
import datetime as dt
import re
import sys

import openpyxl

COLUMNS = ['name', 'company', 'customer_type', 'phone', 'email', 'facebook', 'source', 'need', 'region', 'headcount',
           'event_time', 'stage', 'lost_reason', 'assignee_name', 'next_followup', 'notes', 'created_at']
REGIONS = {'hà nội': 'Hà Nội', 'hải phòng': 'Hải Phòng', 'đà nẵng': 'Đà Nẵng', 'tp.hcm': 'TP.HCM', 'hcm': 'TP.HCM'}


def txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def fix_d(s):
    # File cũ dùng "Ð" (U+00D0, chữ Iceland) lẫn với "Đ" (U+0110) trong dropdown trạng thái
    return s.replace('Ð', 'Đ')


def phone_of(v):
    s = re.sub(r'[^\d]', '', txt(v))
    if len(s) == 9:
        s = '0' + s
    return s if 9 <= len(s) <= 11 else ''


def date_iso(v):
    return v.strftime('%Y-%m-%dT09:00:00') if isinstance(v, (dt.datetime, dt.date)) else ''


def need_of(v):
    s = txt(v).lower()
    if 'pick' in s:
        return 'Pickleball'
    if 'bóng' in s:
        return 'Bóng đá'
    if 'cầu lông' in s:
        return 'Cầu lông'
    if 'hội thao' in s:
        return 'Hội thao'
    return txt(v)


def fb_leads(ws):
    stage_map = {'Chưa liên hệ': 'new', 'Đã liên hệ': 'contacted', 'Đang tư vấn': 'consulting', 'Đã chốt': 'won',
                 'Đã liên hệ nhưng không rep': 'contacted'}
    seen, out = set(), []
    for r in ws.iter_rows(min_row=2, values_only=True):
        _, name, date, channel, region, need, pic, status, note = (list(r) + [None] * 9)[:9]
        name = txt(name)
        if not name or name == 'Tên khách hàng':
            continue
        phone = phone_of(note)
        key = (name.lower(), phone)
        if key in seen:
            continue
        seen.add(key)
        status = fix_d(txt(status))
        notes = [f'Kênh: {txt(channel)}'] if txt(channel) else []
        reg = REGIONS.get(txt(region).lower(), '')
        if txt(region) and not reg:
            notes.append(txt(region))
        if txt(note) and not phone:
            notes.append(txt(note))
        if status == 'Đã liên hệ nhưng không rep':
            notes.append('Đã liên hệ nhưng khách không phản hồi')
        notes.append('[Nhập từ Excel cũ – sheet DS KH ads FB]')
        out.append({
            'name': name, 'customer_type': 'Cá nhân / CLB', 'phone': phone, 'source': 'Facebook Ads (tin nhắn)',
            'need': need_of(need), 'region': reg, 'stage': stage_map.get(status, 'contacted'),
            'assignee_name': txt(pic), 'notes': '\n'.join(notes), 'created_at': date_iso(date),
        })
    return out


def deals(ws):
    stage_map = {'Đang đàm phán': 'negotiating', 'Đã hoàn thành': 'won', 'Đang thực hiện': 'won',
                 'Chưa bắt đầu': 'consulting', 'Bị chặn': 'negotiating'}
    out = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        name, status, owner, phase, pic, due, sport, note, plan, extra = (list(r) + [None] * 10)[:10]
        name = txt(name)
        if not name:
            continue
        status = fix_d(txt(status))
        notes = []
        if status == 'Bị chặn':
            notes.append('⚠ Trạng thái cũ: Bị chặn / tạm dừng')
        for label, v in (('Ghi chú', note), ('Phương án', plan), ('Khác', extra)):
            v = txt(v) if not isinstance(v, dt.datetime) else v.strftime('%d/%m/%Y')
            if v:
                notes.append(f'{label}: {v}')
        notes.append('[Nhập từ Excel cũ – sheet DS Giải đấu]')
        out.append({
            'name': name, 'company': name, 'customer_type': 'Doanh nghiệp', 'source': 'Khác',
            'need': need_of(sport),
            'event_time': due.strftime('%d/%m/%Y') if isinstance(due, dt.datetime) else txt(due),
            'stage': stage_map.get(status, 'consulting'), 'assignee_name': txt(pic) or txt(owner),
            'notes': '\n'.join(notes), 'created_at': date_iso(due) if isinstance(due, dt.datetime) and due < dt.datetime.now() else '',
        })
    return out


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    rows = fb_leads(wb['DS KH ads FB']) + deals(wb['DS Giải đấu'])
    with open(dst, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, '') for k in COLUMNS})
    names = sorted({r['assignee_name'] for r in rows if r['assignee_name']})
    print(f'Đã xuất {len(rows)} lead → {dst}')
    print('Tên nhân viên cần khớp trong CRM:', ', '.join(names))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
